"""Document uploads end-to-end (PDF + Word + text) — highest breakage-risk paths.

Contract locked here:
- POST /api/sources/upload accepts .pdf/.docx/.txt/.md, stores the real
  source_type, and schedules process_source_file (never process_pdf directly).
- Anything else (.exe, no extension) is rejected with 400 naming allowed types.
- Oversize bodies are rejected with 400. Viewers (role=viewer) get 403.
- process_source_file routes by type: pdf -> PyPDFLoader pipeline (unchanged),
  docx/txt/md -> text extraction + shared chunk pipeline. Blank files and
  corrupt files end in status=error, never status=ready with zero chunks.
- GET /api/sources/{id}/download serves the correct media type per file kind.
"""
import io

import pytest
from bson import ObjectId
from fastapi.testclient import TestClient
from unittest.mock import patch, AsyncMock

from app.main import app
from app.routes.sources import settings as sources_settings
from app.services.auth_service import create_access_token
from app.services.rag_service import (
    detect_source_type,
    media_type_for_source_type,
    extract_docx_text,
)

client = TestClient(app)

USER_ID = str(ObjectId())
USER_EMAIL = "docsuser@example.com"
USER_TOKEN = create_access_token(USER_ID, USER_EMAIL)
WORKSPACE_ID = str(ObjectId())
SOURCE_ID = str(ObjectId())


def _auth_headers():
    return {"Authorization": f"Bearer {USER_TOKEN}"}


def _mock_user(mock_users):
    mock_users.find_one = AsyncMock(return_value={
        "_id": ObjectId(USER_ID),
        "email": USER_EMAIL,
        "username": "Docs User",
        "active_workspace_id": WORKSPACE_ID,
    })


def _mock_workspace_member(mock_workspaces, role="owner"):
    mock_workspaces.find_one = AsyncMock(return_value={
        "_id": ObjectId(WORKSPACE_ID),
        "owner_id": USER_ID,
        "name": "HQ",
        "members": [{"user_id": USER_ID, "email": USER_EMAIL, "role": role}],
    })


def _mock_workspace_viewer(mock_workspaces):
    # owner_id differs so role_of() resolves the caller as viewer, not owner.
    mock_workspaces.find_one = AsyncMock(return_value={
        "_id": ObjectId(WORKSPACE_ID),
        "owner_id": str(ObjectId()),
        "name": "HQ",
        "members": [{"user_id": USER_ID, "email": USER_EMAIL, "role": "viewer"}],
    })


def _mocked_db(mock_users, mock_sources, mock_workspaces):
    _mock_user(mock_users)
    _mock_workspace_member(mock_workspaces)
    mock_sources.insert_one = AsyncMock(
        return_value=type("R", (), {"inserted_id": ObjectId(SOURCE_ID)})()
    )


def _upload(filename, content, content_type):
    with patch("app.database.users_collection") as mock_users, \
         patch("app.database.sources_collection") as mock_sources, \
         patch("app.database.workspaces_collection") as mock_workspaces, \
         patch("app.routes.sources.check_db"), \
         patch("app.services.storage.is_r2_enabled", return_value=True), \
         patch("app.services.storage.upload_file_bytes",
               return_value=f"{USER_ID}/{SOURCE_ID}_{filename}"), \
         patch("app.routes.sources.process_source_file") as mock_process:
        _mocked_db(mock_users, mock_sources, mock_workspaces)
        response = client.post(
            "/api/sources/upload",
            files={"file": (filename, io.BytesIO(content), content_type)},
            headers=_auth_headers(),
        )
        return response, mock_process


@pytest.mark.parametrize(
    "filename,content_type,expected_type",
    [
        ("report.pdf", "application/pdf", "pdf"),
        ("report.docx",
         "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
         "docx"),
        ("notes.txt", "text/plain", "txt"),
        ("notes.md", "text/markdown", "md"),
        ("REPORT.DOCX",
         "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
         "docx"),
    ],
)
def test_upload_accepts_pdf_word_and_text(filename, content_type, expected_type):
    response, mock_process = _upload(filename, b"fake-bytes", content_type)
    assert response.status_code == 201
    data = response.json()
    assert data["source_type"] == expected_type
    assert data["status"] == "processing"
    assert mock_process.called


def test_upload_rejects_executable():
    with patch("app.database.users_collection") as mock_users, \
         patch("app.database.sources_collection") as mock_sources, \
         patch("app.database.workspaces_collection") as mock_workspaces, \
         patch("app.routes.sources.check_db"):
        _mocked_db(mock_users, mock_sources, mock_workspaces)
        response = client.post(
            "/api/sources/upload",
            files={"file": ("evil.exe", io.BytesIO(b"MZ"), "application/octet-stream")},
            headers=_auth_headers(),
        )
    assert response.status_code == 400
    assert ".pdf" in response.json()["detail"]
    assert ".docx" in response.json()["detail"]


def test_upload_rejects_extensionless_filename():
    with patch("app.database.users_collection") as mock_users, \
         patch("app.database.sources_collection") as mock_sources, \
         patch("app.database.workspaces_collection") as mock_workspaces, \
         patch("app.routes.sources.check_db"):
        _mocked_db(mock_users, mock_sources, mock_workspaces)
        response = client.post(
            "/api/sources/upload",
            files={"file": ("README", io.BytesIO(b"hi"), "text/plain")},
            headers=_auth_headers(),
        )
    assert response.status_code == 400


def test_upload_rejects_oversize_body():
    with patch("app.database.users_collection") as mock_users, \
         patch("app.database.sources_collection") as mock_sources, \
         patch("app.database.workspaces_collection") as mock_workspaces, \
         patch("app.routes.sources.check_db"), \
         patch.object(sources_settings, "MAX_FILE_SIZE_MB", 0):
        _mocked_db(mock_users, mock_sources, mock_workspaces)
        response = client.post(
            "/api/sources/upload",
            files={"file": ("big.pdf", io.BytesIO(b"12345"), "application/pdf")},
            headers=_auth_headers(),
        )
    assert response.status_code == 400
    assert "exceeds" in response.json()["detail"]


def test_upload_viewer_forbidden_for_docx():
    with patch("app.database.users_collection") as mock_users, \
         patch("app.database.sources_collection") as mock_sources, \
         patch("app.database.workspaces_collection") as mock_workspaces, \
         patch("app.routes.sources.check_db"):
        _mock_user(mock_users)
        _mock_workspace_viewer(mock_workspaces)
        mock_sources.insert_one = AsyncMock(
            return_value=type("R", (), {"inserted_id": ObjectId(SOURCE_ID)})()
        )
        response = client.post(
            "/api/sources/upload",
            files={"file": ("report.docx",
                            io.BytesIO(b"fake"),
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers=_auth_headers(),
        )
    assert response.status_code == 403


def test_detect_source_type():
    assert detect_source_type("a.pdf") == "pdf"
    assert detect_source_type("Report.DOCX") == "docx"
    assert detect_source_type("notes.txt") == "txt"
    assert detect_source_type("notes.MD") == "md"
    assert detect_source_type("evil.exe") is None
    assert detect_source_type("no-extension") is None
    assert detect_source_type("") is None


def test_media_types():
    assert media_type_for_source_type("pdf") == "application/pdf"
    assert media_type_for_source_type("docx").startswith(
        "application/vnd.openxmlformats")
    assert media_type_for_source_type("txt") == "text/plain"
    assert media_type_for_source_type("md") == "text/markdown"
    assert media_type_for_source_type("unknown") == "application/octet-stream"


def test_extract_docx_text_with_table(tmp_path):
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("Hello Quello")
    doc.add_paragraph("")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "cell-one"
    table.cell(0, 1).text = "cell-two"
    path = tmp_path / "sample.docx"
    doc.save(str(path))

    text = extract_docx_text(str(path))
    assert "Hello Quello" in text
    assert "cell-one" in text
    assert "cell-two" in text


def test_extract_docx_text_rejects_corrupt(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip file")
    with pytest.raises(ValueError):
        extract_docx_text(str(bad))


@pytest.mark.asyncio
async def test_process_blank_txt_marks_error(tmp_path):
    from app.services import rag_service

    empty = tmp_path / "empty.txt"
    empty.write_bytes(b"   \n  ")
    with patch("app.database.sources_collection") as mock_sources:
        mock_sources.update_one = AsyncMock()
        with pytest.raises(ValueError):
            await rag_service.process_source_file(
                USER_ID, SOURCE_ID, str(empty), WORKSPACE_ID,
                source_type="txt",
            )
        _filter, update = mock_sources.update_one.call_args[0]
        assert update["$set"]["status"] == "error"


@pytest.mark.asyncio
async def test_process_txt_indexes_and_marks_ready(tmp_path):
    from app.services import rag_service

    hello = tmp_path / "hello.txt"
    hello.write_bytes(b"Quello indexes plain text files for grounded answers.")
    with patch("app.database.sources_collection") as mock_sources, \
         patch.object(rag_service.vs, "delete_source_vectors", return_value=0), \
         patch.object(rag_service.vs, "add_documents", return_value=3):
        mock_sources.update_one = AsyncMock()
        count = await rag_service.process_source_file(
            USER_ID, SOURCE_ID, str(hello), WORKSPACE_ID,
            source_type="txt",
        )
        assert count == 3
        _filter, update = mock_sources.update_one.call_args[0]
        assert update["$set"]["status"] == "ready"
        assert update["$set"]["chunk_count"] == 3


def test_download_docx_serves_word_mime_type(tmp_path):
    local_file = tmp_path / "report.docx"
    local_file.write_bytes(b"PK fake docx")
    with patch("app.database.users_collection") as mock_users, \
         patch("app.database.sources_collection") as mock_sources, \
         patch("app.database.workspaces_collection") as mock_workspaces, \
         patch("app.routes.sources.check_db"):
        _mock_user(mock_users)
        _mock_workspace_member(mock_workspaces)
        mock_sources.find_one = AsyncMock(return_value={
            "_id": ObjectId(SOURCE_ID),
            "user_id": USER_ID,
            "workspace_id": WORKSPACE_ID,
            "filename": "report.docx",
            "file_size": 12,
            "status": "ready",
            "source_type": "docx",
            "file_path": str(local_file),
        })
        response = client.get(
            f"/api/sources/{SOURCE_ID}/download",
            headers=_auth_headers(),
        )
    assert response.status_code == 200
    assert response.headers.get("content-type", "").startswith(
        "application/vnd.openxmlformats")
